from io import BytesIO
from openpyxl import load_workbook

def parse_skus_from_xlsx(content: bytes):
    wb = load_workbook(filename=BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    skus = []
    for row in ws.iter_rows(min_row=1, max_col=1, values_only=True):
        val = (row[0] or "").strip() if row[0] else ""
        if val:
            skus.append(val)
    return skus
