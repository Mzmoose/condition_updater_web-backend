import io, re, csv
from typing import List
from openpyxl import load_workbook

FIRST4_RE = re.compile(r"^\s*(\d{4})")

def first4(s: str) -> str:
    if not s:
        return ""
    m = FIRST4_RE.match(str(s))
    return m.group(1) if m else ""

def parse_file_to_skus(binary: bytes, filename: str) -> List[str]:
    def push(acc: List[str], raw: str):
        key = first4(raw)
        if key:
            acc.append(key)

    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        wb = load_workbook(io.BytesIO(binary), read_only=True, data_only=True)
        skus: List[str] = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                if not row:
                    continue
                v = row[0]
                if v is None:
                    continue
                push(skus, str(v))
        seen, out = set(), []
        for s in skus:
            if s not in seen:
                seen.add(s)
                out.append(s)
        return out

    if name.endswith(".csv"):
        text = binary.decode("utf-8", errors="replace")
        rdr = csv.reader(io.StringIO(text))
        skus: List[str] = []
        for i, row in enumerate(rdr):
            if not row:
                continue
            raw = str(row[0])
            if i == 0 and raw.strip().lower() in {"sku", "customlabel"} and len(row) == 1:
                continue
            push(skus, raw)
        return skus

    text = binary.decode("utf-8", errors="replace")
    skus: List[str] = []
    for ln in text.splitlines():
        push(skus, ln)
    return [s for s in skus if s]
