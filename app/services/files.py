from __future__ import annotations
from io import BytesIO
from typing import List
import csv, re

try:
    from openpyxl import load_workbook  # for .xlsx
except Exception:  # pragma: no cover
    load_workbook = None  # type: ignore

FOUR_DIGIT = re.compile(r"^(\d{4})")
TOKEN = re.compile(r"^([A-Za-z0-9\-]+)")

def _norm(raw: str) -> str:
    s = (raw or "").strip()
    if not s:
        return ""
    m = FOUR_DIGIT.match(s)
    if m:
        return m.group(1)
    m = TOKEN.match(s)
    return m.group(1) if m else s

def parse_file_to_skus(file_bytes: bytes, filename: str | None = None) -> List[str]:
    """
    Accept .xlsx, .csv, or .txt (one SKU per line).
    Return a de-duplicated list of SKUs, normalized to the first 4 digits when present.
    """
    name = (filename or "").lower()
    out: list[str] = []

    def push(v) -> None:
        val = _norm(str(v))
        if val and val not in out:
            out.append(val)

    # Excel
    if name.endswith((".xlsx", ".xlsm", ".xltx", ".xltm")):
        if load_workbook is None:
            raise RuntimeError("openpyxl is required to read .xlsx files")
        wb = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
        ws = wb.active
        for row in ws.iter_rows(min_row=1, values_only=True):
            if row and row[0] is not None:
                push(row[0])
        return out

    # CSV
    if name.endswith(".csv"):
        text = BytesIO(file_bytes).read().decode("utf-8", errors="ignore")
        for row in csv.reader(text.splitlines()):
            if row:
                push(row[0])
        return out

    # Plain text
    text = BytesIO(file_bytes).read().decode("utf-8", errors="ignore")
    for line in text.splitlines():
        if line.strip():
            push(line.strip())
    return out
