from fastapi import APIRouter, UploadFile, File, HTTPException
import io
from openpyxl import load_workbook

router = APIRouter()

def _to_str(x):
    return "" if x is None else str(x).strip()

@router.post("/condition/update")
async def condition_update(file: UploadFile = File(...)):
    try:
        data = await file.read()
        wb = load_workbook(io.BytesIO(data), data_only=True)
        ws = wb.active

        rows = []
        # No headers: treat row 1 as data. Col1 = SKU, Col2 = Condition
        for r in range(1, ws.max_row + 1):
            sku  = _to_str(ws.cell(row=r, column=1).value)
            cond = _to_str(ws.cell(row=r, column=2).value)
            if not sku or not cond:
                continue
            rows.append({"sku": sku, "condition": cond})

        if not rows:
            raise HTTPException(
                status_code=400,
                detail="No valid rows found (expected 2 columns: SKU, Condition)."
            )

        # For now return a preview so we can confirm parsing works:
        return {"ok": True, "parsed": len(rows), "sample": rows[:5]}

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to read Excel: {e}")
